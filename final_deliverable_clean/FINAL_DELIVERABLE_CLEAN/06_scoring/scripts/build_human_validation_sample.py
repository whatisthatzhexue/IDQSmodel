from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from scoring_utils import SCORING_ROOT, read_csv_rows, save_json, write_csv_rows
from stability_common import safe_float, write_markdown


SAMPLE_HEADERS = [
    "sample_id",
    "sample_reason",
    "document_id",
    "ticker",
    "company_name",
    "report_year",
    "total_score_100",
    "grade_label",
    "dimension_code",
    "dimension_name",
    "model_raw_score_1_5",
    "model_confidence_level",
    "evidence_locator",
    "evidence_text_short",
    "model_comment_short",
    "numeric_check_path",
    "evidence_semantic_status",
    "evidence_issue_type",
    "score_evidence_alignment",
    "recommended_action",
    "human_score_1_5",
    "human_evidence_valid_y_n",
    "human_dimension_fit_y_n",
    "human_notes",
    "reviewer_id",
]

DIMENSION_ORDER = ["AR01", "AR02", "AR03", "AR04", "AR05"]
REASON_ORDER = ["semantic_review", "ar02_numeric_review", "low_score", "high_score", "baseline"]


def build_human_validation_sample(root: Path = SCORING_ROOT, sample_size: int = 30, v2: bool = False) -> dict[str, Any]:
    out_dir = root / "HUMAN_VALIDATION"
    out_dir.mkdir(parents=True, exist_ok=True)
    if v2:
        docs = _preferred_rows(root / "06_ratings" / "mda_final_v2" / "mda_final_document_scores.csv", root / "06_ratings" / "mda_final_full" / "mda_final_document_scores.csv")
        ratings = _preferred_rows(root / "06_ratings" / "mda_final_v2" / "mda_final_ratings_long.csv", root / "06_ratings" / "mda_final_full" / "mda_final_ratings_long.csv")
        audit_rows = _preferred_rows(root / "08_reports" / "mda_semantic_evidence_audit_v2.csv", root / "08_reports" / "mda_semantic_evidence_audit.csv")
    else:
        docs = _preferred_rows(
            root / "06_ratings" / "mda_final_full" / "mda_final_document_scores_repaired.csv",
            root / "06_ratings" / "mda_final_full" / "mda_final_document_scores.csv",
        )
        ratings = _preferred_rows(
            root / "06_ratings" / "mda_final_full" / "mda_final_ratings_long_repaired.csv",
            root / "06_ratings" / "mda_final_full" / "mda_final_ratings_long.csv",
        )
        audit_rows = read_csv_rows(root / "08_reports" / "mda_semantic_evidence_audit.csv")
    doc_by_id = {row.get("document_id", ""): row for row in docs}
    audit_by_key = {(row.get("document_id", ""), row.get("dimension_code", "")): row for row in audit_rows}

    candidates = [_sample_candidate(row, doc_by_id.get(row.get("document_id", ""), {}), audit_by_key.get((row.get("document_id", ""), row.get("dimension_code", "")), {})) for row in ratings]
    candidates = [row for row in candidates if row.get("document_id") and row.get("dimension_code")]
    sample_rows = _select_v2_sample(candidates, sample_size) if v2 else _select_sample(candidates, sample_size)
    for idx, row in enumerate(sample_rows, start=1):
        row["sample_id"] = f"HVAL_{idx:03d}"

    suffix = "_v2" if v2 else ""
    csv_path = out_dir / f"mda_human_validation_sample{suffix}.csv"
    xlsx_path = out_dir / f"mda_human_validation_sample{suffix}.xlsx"
    instructions_path = out_dir / f"mda_human_validation_instructions{suffix}.md"
    summary_path = out_dir / f"human_validation_summary{suffix}.json"
    write_csv_rows(csv_path, SAMPLE_HEADERS, sample_rows)
    _write_xlsx(xlsx_path, sample_rows)
    _write_instructions(instructions_path, sample_size, len(sample_rows), v2=v2)
    summary = {
        "sample_rows": len(sample_rows),
        "unique_document_count": len({row.get("document_id", "") for row in sample_rows}),
        "requested_sample_size": sample_size,
        "source_rating_rows": len(ratings),
        "source_document_rows": len(docs),
        "mda_human_validation_ready": False,
        "xlsx_path": str(xlsx_path),
        "csv_path": str(csv_path),
        "instructions_path": str(instructions_path),
    }
    save_json(summary_path, summary)
    return summary


def _preferred_rows(primary: Path, fallback: Path) -> list[dict[str, str]]:
    primary_rows = read_csv_rows(primary)
    if primary_rows:
        return primary_rows
    return read_csv_rows(fallback)


def _sample_candidate(rating: dict[str, str], doc: dict[str, str], audit: dict[str, str]) -> dict[str, Any]:
    raw_score = safe_float(rating.get("raw_score", ""), 0.0)
    dimension = rating.get("dimension_code", "")
    semantic_status = audit.get("evidence_semantic_status", "")
    recommended_action = audit.get("recommended_action", "")
    needs_review = audit.get("needs_review", "").lower() in {"1", "true", "yes"}
    if semantic_status in {"invalid", "weak"} or needs_review or recommended_action not in {"", "accept"}:
        reason = "semantic_review"
    elif dimension == "AR02" and rating.get("numeric_check_path"):
        reason = "ar02_numeric_review"
    elif raw_score <= 2:
        reason = "low_score"
    elif raw_score >= 4:
        reason = "high_score"
    else:
        reason = "baseline"
    return {
        "sample_id": "",
        "sample_reason": reason,
        "document_id": rating.get("document_id", ""),
        "ticker": doc.get("ticker", ""),
        "company_name": doc.get("company_name", ""),
        "report_year": doc.get("report_year", ""),
        "total_score_100": doc.get("total_score_100", ""),
        "grade_label": doc.get("grade_label", ""),
        "dimension_code": dimension,
        "dimension_name": rating.get("dimension_name", ""),
        "model_raw_score_1_5": rating.get("raw_score", ""),
        "model_confidence_level": rating.get("confidence_level", ""),
        "evidence_locator": rating.get("evidence_locator", "") or audit.get("evidence_locator", ""),
        "evidence_text_short": audit.get("evidence_text_short", "") or rating.get("evidence_text_short", ""),
        "model_comment_short": rating.get("comment_short", ""),
        "numeric_check_path": rating.get("numeric_check_path", ""),
        "evidence_semantic_status": semantic_status,
        "evidence_issue_type": audit.get("evidence_issue_type", ""),
        "score_evidence_alignment": audit.get("score_evidence_alignment", ""),
        "recommended_action": recommended_action,
        "human_score_1_5": "",
        "human_evidence_valid_y_n": "",
        "human_dimension_fit_y_n": "",
        "human_notes": "",
        "reviewer_id": "",
    }


def _select_sample(candidates: list[dict[str, Any]], sample_size: int) -> list[dict[str, Any]]:
    selected: list[dict[str, Any]] = []
    seen: set[tuple[str, str]] = set()

    def add(row: dict[str, Any]) -> None:
        key = (row.get("document_id", ""), row.get("dimension_code", ""))
        if key in seen or len(selected) >= sample_size:
            return
        selected.append(dict(row))
        seen.add(key)

    for reason in REASON_ORDER:
        for dimension in DIMENSION_ORDER:
            rows = [row for row in candidates if row["sample_reason"] == reason and row["dimension_code"] == dimension]
            if rows:
                add(sorted(rows, key=_sort_key)[0])
    for row in sorted(candidates, key=_sort_key):
        add(row)
        if len(selected) >= sample_size:
            break
    return selected


def _select_v2_sample(candidates: list[dict[str, Any]], document_count: int) -> list[dict[str, Any]]:
    by_doc: dict[str, list[dict[str, Any]]] = {}
    for row in candidates:
        by_doc.setdefault(row["document_id"], []).append(row)
    prioritized_docs = []
    for doc_id, rows in by_doc.items():
        reason_rank = min(REASON_ORDER.index(row["sample_reason"]) if row["sample_reason"] in REASON_ORDER else len(REASON_ORDER) for row in rows)
        scores = [safe_float(row.get("total_score_100", ""), 0.0) for row in rows]
        total = scores[0] if scores else 0.0
        prioritized_docs.append((reason_rank, _score_bucket_rank(total), doc_id))
    selected_docs = [doc_id for _, _, doc_id in sorted(prioritized_docs)[:document_count]]
    out = []
    for doc_id in selected_docs:
        rows = {row["dimension_code"]: row for row in by_doc.get(doc_id, [])}
        for code in DIMENSION_ORDER:
            if code in rows:
                out.append(dict(rows[code]))
    return out


def _score_bucket_rank(total: float) -> int:
    if total < 55:
        return 0
    if total >= 85:
        return 1
    if total >= 70:
        return 2
    return 3


def _sort_key(row: dict[str, Any]) -> tuple[int, int, str, str]:
    reason_rank = REASON_ORDER.index(row["sample_reason"]) if row["sample_reason"] in REASON_ORDER else len(REASON_ORDER)
    dimension_rank = DIMENSION_ORDER.index(row["dimension_code"]) if row["dimension_code"] in DIMENSION_ORDER else len(DIMENSION_ORDER)
    return (reason_rank, dimension_rank, row.get("document_id", ""), row.get("dimension_code", ""))


def _write_xlsx(path: Path, rows: list[dict[str, Any]]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "MD&A Validation"
    ws.append(SAMPLE_HEADERS)
    for row in rows:
        ws.append([row.get(header, "") for header in SAMPLE_HEADERS])
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    human_fill = PatternFill("solid", fgColor="FFF2CC")
    for col_idx, header in enumerate(SAMPLE_HEADERS, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        cell.fill = human_fill if header.startswith("human_") or header == "reviewer_id" else header_fill
        ws.column_dimensions[get_column_letter(col_idx)].width = 18
    for column in ["N", "O", "X"]:
        ws.column_dimensions[column].width = 46
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def _write_instructions(path: Path, requested: int, actual: int, *, v2: bool = False) -> None:
    lines = [
        "# MD&A Human Validation Instructions",
        "",
        f"- Requested sample size: {requested}",
        f"- Actual sample rows: {actual}",
        f"- Sampling mode: {'30 unique MD&A documents x 5 dimensions' if v2 else 'dimension-row sample'}",
        "- Scope: MD&A-only model-generated candidate dataset.",
        "- Do not edit the model score columns.",
        "- Fill only the human_score_1_5, human_evidence_valid_y_n, human_dimension_fit_y_n, human_notes, and reviewer_id columns.",
        "- Use human_score_1_5 only when the evidence supports an independent 1-5 judgment for the listed dimension.",
        "- Mark human_evidence_valid_y_n as N if the locator or excerpt does not support the dimension.",
        "- Pay extra attention to sample_reason values semantic_review and ar02_numeric_review.",
        "- News scoring is outside this MD&A human-validation sample; use the News review queues for News-specific manual checks.",
    ]
    write_markdown(path, lines)


def main() -> int:
    parser = argparse.ArgumentParser(description="Build a deterministic MD&A human-validation sample workbook.")
    parser.add_argument("--root", type=Path, default=SCORING_ROOT)
    parser.add_argument("--sample-size", type=int, default=30)
    parser.add_argument("--v2", action="store_true")
    args = parser.parse_args()
    print(json.dumps(build_human_validation_sample(args.root, args.sample_size, args.v2), indent=2, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
