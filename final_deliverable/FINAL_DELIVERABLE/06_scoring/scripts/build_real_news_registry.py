from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from scoring_utils import SCORING_ROOT, write_csv_rows
from stability_common import write_markdown


REQUIRED_FIELDS = ["ticker", "company_name", "publish_date", "source_name", "title", "url", "article_text"]
REAL_NEWS_REGISTRY_HEADERS = [
    "document_id",
    "doc_type",
    "ticker",
    "company_name",
    "publish_date",
    "source_name",
    "title",
    "url",
    "file_path",
    "article_text_hash",
    "text_length",
    "synthetic_flag",
    "include_flag",
    "exclude_reason",
    "created_at",
]


def build_real_news_registry(root: Path = SCORING_ROOT, source_path: Path | None = None) -> dict[str, Any]:
    source = source_path or _discover_source(root)
    if source is None:
        summary = _empty_summary("missing_real_news_input")
        _write_registry_report(root, summary)
        _write_missing_blocker(root)
        return summary

    rows, field_error = _read_source_rows(source)
    if field_error:
        summary = _empty_summary(field_error)
        summary["source_path"] = str(source)
        _write_registry_report(root, summary)
        _write_missing_blocker(root)
        return summary

    created_at = datetime.now(timezone.utc).replace(microsecond=0).isoformat()
    registry_rows: list[dict[str, str]] = []
    seen_article_hashes: set[str] = set()
    seen_title_keys: set[tuple[str, str, str, str]] = set()
    counts = {
        "duplicate_count": 0,
        "missing_article_text_count": 0,
        "short_text_count": 0,
        "invalid_date_count": 0,
        "synthetic_count": 0,
    }

    for source_index, row in enumerate(rows, start=1):
        ticker = _clean(row.get("ticker", "")).upper()
        company_name = _clean(row.get("company_name", ""))
        publish_date_raw = _clean(row.get("publish_date", ""))
        publish_date = _parse_date(publish_date_raw)
        source_name = _clean(row.get("source_name", ""))
        title = _clean(row.get("title", ""))
        url = _clean(row.get("url", ""))
        article_text = str(row.get("article_text", "") or "")
        normalized_text = _normalize_article_text(article_text)
        compact_len = len(re.sub(r"\s+", "", normalized_text))
        article_hash = hashlib.sha256(normalized_text.encode("utf-8")).hexdigest()
        title_key = (ticker, source_name.lower(), publish_date or publish_date_raw, title.lower())

        reasons = []
        if not normalized_text:
            reasons.append("missing_article_text")
            counts["missing_article_text_count"] += 1
        if compact_len < 200:
            reasons.append("short_article_text")
            counts["short_text_count"] += 1
        if not publish_date:
            reasons.append("invalid_publish_date")
            counts["invalid_date_count"] += 1
        if _looks_placeholder(normalized_text, title, url):
            reasons.append("placeholder_or_non_article_text")
        if article_hash in seen_article_hashes or title_key in seen_title_keys:
            reasons = ["duplicate"]
            counts["duplicate_count"] += 1

        include_flag = "No" if reasons else "Yes"
        if include_flag == "Yes":
            seen_article_hashes.add(article_hash)
            seen_title_keys.add(title_key)

        doc_date = (publish_date or publish_date_raw or "unknown").replace("-", "")
        doc_id = _document_id(ticker or "NEWS", doc_date, article_hash, source_index)
        registry_rows.append(
            {
                "document_id": doc_id,
                "doc_type": "news",
                "ticker": ticker,
                "company_name": company_name,
                "publish_date": publish_date or publish_date_raw,
                "source_name": source_name,
                "title": title,
                "url": url,
                "file_path": _relative_or_absolute(root, source),
                "article_text_hash": article_hash,
                "text_length": str(compact_len),
                "synthetic_flag": "false",
                "include_flag": include_flag,
                "exclude_reason": ";".join(reasons),
                "created_at": created_at,
            }
        )

    write_csv_rows(root / "01_registry" / "news_registry.csv", REAL_NEWS_REGISTRY_HEADERS, registry_rows)
    included = sum(row["include_flag"] == "Yes" for row in registry_rows)
    summary = {
        "status": "success" if included else "blocked",
        "source_path": str(source),
        "registry_path": str(root / "01_registry" / "news_registry.csv"),
        "total_news_rows": len(registry_rows),
        "included_news_rows": included,
        "excluded_news_rows": len(registry_rows) - included,
        **counts,
    }
    _write_registry_report(root, summary)
    if included == 0:
        _write_missing_blocker(root)
    return summary


def _discover_source(root: Path) -> Path | None:
    input_dir = root / "input_news"
    preferred = [input_dir / "news_raw.xlsx", input_dir / "news_raw.csv"]
    for path in preferred:
        if path.exists():
            return path
    for pattern in ("*.xlsx", "*.csv"):
        matches = sorted(path for path in input_dir.glob(pattern) if path.is_file())
        if matches:
            return matches[0]
    return None


def _read_source_rows(path: Path) -> tuple[list[dict[str, str]], str | None]:
    if path.suffix.lower() == ".csv":
        with path.open(newline="", encoding="utf-8-sig") as fh:
            reader = csv.DictReader(fh)
            headers = set(reader.fieldnames or [])
            missing = [field for field in REQUIRED_FIELDS if field not in headers]
            if missing:
                return [], "missing_required_fields:" + ",".join(missing)
            return [dict(row) for row in reader], None
    if path.suffix.lower() == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError:
            return [], "openpyxl_not_installed_for_xlsx"
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return [], None
        headers = [str(value or "").strip() for value in rows[0]]
        missing = [field for field in REQUIRED_FIELDS if field not in headers]
        if missing:
            return [], "missing_required_fields:" + ",".join(missing)
        out = []
        for values in rows[1:]:
            out.append({headers[idx]: "" if value is None else str(value) for idx, value in enumerate(values[: len(headers)])})
        return out, None
    return [], "unsupported_news_input_format"


def _parse_date(value: str) -> str:
    text = _clean(value)
    if not text:
        return ""
    for pattern in [r"(\d{4})[-/](\d{1,2})[-/](\d{1,2})", r"(\d{1,2})[-/](\d{1,2})[-/](\d{4})"]:
        match = re.search(pattern, text)
        if not match:
            continue
        parts = match.groups()
        if len(parts[0]) == 4:
            year, month, day = parts
        else:
            day, month, year = parts
        try:
            parsed = datetime(int(year), int(month), int(day))
        except ValueError:
            return ""
        return parsed.strftime("%Y-%m-%d")
    return ""


def _normalize_article_text(value: str) -> str:
    return "\n".join(line.strip() for line in str(value or "").replace("\r", "\n").split("\n")).strip()


def _looks_placeholder(text: str, title: str, url: str) -> bool:
    compact = re.sub(r"\s+", "", text).lower()
    if not compact:
        return False
    lowered = text.strip().lower()
    if lowered in {title.strip().lower(), url.strip().lower()}:
        return True
    if re.fullmatch(r"https?://\S+", lowered):
        return True
    placeholder_terms = ["placeholder", "lorem ipsum", "search result", "click here", "read more"]
    return any(term in lowered for term in placeholder_terms)


def _document_id(ticker: str, publish_date_compact: str, article_hash: str, source_index: int) -> str:
    digest = hashlib.sha1(f"{ticker}|{publish_date_compact}|{article_hash}|{source_index}".encode("utf-8")).hexdigest()[:8].upper()
    safe_ticker = re.sub(r"[^A-Z0-9]", "", ticker.upper()) or "NEWS"
    safe_date = re.sub(r"[^0-9A-Za-z]", "", publish_date_compact) or "UNKNOWN"
    return f"NEWS_{safe_ticker}_{safe_date}_{digest}"


def _clean(value: Any) -> str:
    return " ".join(str(value or "").replace("\r", "\n").split())


def _relative_or_absolute(root: Path, path: Path) -> str:
    try:
        return str(path.relative_to(root))
    except ValueError:
        return str(path)


def _empty_summary(reason: str) -> dict[str, Any]:
    return {
        "status": "blocked",
        "reason": reason,
        "total_news_rows": 0,
        "included_news_rows": 0,
        "excluded_news_rows": 0,
        "duplicate_count": 0,
        "missing_article_text_count": 0,
        "short_text_count": 0,
        "invalid_date_count": 0,
        "synthetic_count": 0,
    }


def _write_registry_report(root: Path, summary: dict[str, Any]) -> None:
    lines = ["# Real News Registry Report", ""]
    for key in [
        "status",
        "reason",
        "source_path",
        "total_news_rows",
        "included_news_rows",
        "excluded_news_rows",
        "duplicate_count",
        "missing_article_text_count",
        "short_text_count",
        "invalid_date_count",
        "synthetic_count",
    ]:
        if key in summary:
            lines.append(f"- {key}: {summary[key]}")
    write_markdown(root / "08_reports" / "real_news_registry_report.md", lines)


def _write_missing_blocker(root: Path) -> None:
    write_markdown(
        root / "PAPER_OUTPUT" / "00_status" / "real_news_missing_blocker.md",
        [
            "# Real News Missing Blocker",
            "",
            "MD&A final is complete and frozen.",
            "Combined MD&A + News paper-ready result is blocked because real full-text News data are missing.",
            "",
            "- required_input: 06_scoring/input_news/news_raw.xlsx or 06_scoring/input_news/news_raw.csv",
            "- required_field: article_text",
            "- synthetic_news_allowed: false",
            "- empirical_cross_validation_allowed: false",
        ],
    )


def main() -> int:
    parser = argparse.ArgumentParser(description="Build strict real-News registry from input_news/news_raw.csv or .xlsx.")
    parser.add_argument("--root", type=Path, default=SCORING_ROOT)
    parser.add_argument("--source-path", type=Path)
    args = parser.parse_args()
    print(json.dumps(build_real_news_registry(root=args.root, source_path=args.source_path), indent=2, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
