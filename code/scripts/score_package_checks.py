from __future__ import annotations

import csv
import sys
from pathlib import Path


MDA_WEIGHTS = {"AR01": 0.15, "AR02": 0.25, "AR03": 0.30, "AR04": 0.20, "AR05": 0.10}
BANNED_NAMES = {
    "__MACOSX",
    "__pycache__",
    ".pytest_cache",
    ".venv",
    "archive",
    "_self_correction",
    "mock",
    "benchmark",
    "verification_logs",
    "README.md",
    "PACKAGE_STATUS",
    "PACKAGE_MANIFEST",
    "PAPER_OUTPUT",
    "FINAL_OUTPUT",
}
FORBIDDEN_USER_PATH = "/Users/" + "zhaojiahao"


def read_csv_rows(path: Path) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8") as fh:
        return [dict(row) for row in csv.DictReader(fh)]


def assert_csv_has_header(path: Path) -> None:
    assert path.exists(), f"missing CSV: {path}"
    assert path.stat().st_size > 0, f"0-byte CSV: {path}"
    with path.open(newline="", encoding="utf-8") as fh:
        reader = csv.reader(fh)
        header = next(reader, [])
    assert header, f"missing CSV header: {path}"


def standardize(raw: float) -> float:
    return 100 * (raw - 1) / 4


def check_mda(results_dir: Path) -> None:
    docs = read_csv_rows(results_dir / "MDA" / "mda_final_document_scores.csv")
    ratings = read_csv_rows(results_dir / "MDA" / "mda_final_ratings_long.csv")
    assert len(docs) == 149, f"MDA document rows expected 149, got {len(docs)}"
    assert len(ratings) == 745, f"MDA ratings rows expected 745, got {len(ratings)}"
    for row in docs:
        total = 0.0
        for code, weight in MDA_WEIGHTS.items():
            total += standardize(float(row[code])) * weight
        assert abs(total - float(row["total_score_100"])) < 0.02, f"MDA total mismatch: {row['document_id']}"


def check_human_validation(package_root: Path) -> None:
    xlsx = package_root / "results" / "HUMAN_VALIDATION" / "mda_human_validation_sample_v2.xlsx"
    assert xlsx.exists() and xlsx.stat().st_size > 0, "missing human validation xlsx"
    try:
        from openpyxl import load_workbook  # type: ignore
    except Exception as exc:
        raise AssertionError("openpyxl is required to verify human validation xlsx") from exc
    workbook = load_workbook(xlsx, read_only=True, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    assert rows, "human validation xlsx has no rows"
    header = [str(cell or "") for cell in rows[0]]
    records = rows[1:]
    assert len(records) == 150, f"human validation rows expected 150, got {len(records)}"
    doc_col = header.index("document_id")
    unique_docs = {str(row[doc_col]) for row in records if row[doc_col]}
    assert len(unique_docs) == 30, f"unique human validation documents expected 30, got {len(unique_docs)}"


def check_banned_and_paths(package_root: Path) -> None:
    for path in package_root.rglob("*"):
        parts = set(path.parts)
        assert not (parts & BANNED_NAMES), f"banned path present: {path}"
        assert not path.name.startswith("._"), f"AppleDouble metadata present: {path}"
        assert path.name != ".DS_Store", f".DS_Store present: {path}"
        if path.is_file() and path.suffix.lower() in {".csv", ".md", ".txt", ".json", ".py", ".sh", ".yaml", ".yml"}:
            text = path.read_text(encoding="utf-8", errors="ignore")
            assert FORBIDDEN_USER_PATH not in text, f"absolute user path present: {path}"
    for csv_path in package_root.rglob("*.csv"):
        assert_csv_has_header(csv_path)


def main() -> int:
    code_dir = Path(__file__).resolve().parents[1]
    package_root = code_dir.parent
    results_dir = package_root / "results"
    check_mda(results_dir)
    for path in [
        results_dir / "NEWS" / "news_final_document_scores.csv",
        results_dir / "NEWS" / "news_final_ratings_long.csv",
        results_dir / "NEWS" / "news_final_failed_cases.csv",
        results_dir / "NEWS" / "filter_news_relevance_audit.csv",
        results_dir / "CROSS_VALIDATION" / "eligible_pairs.csv",
        results_dir / "CROSS_VALIDATION" / "claim_links.csv",
        results_dir / "CROSS_VALIDATION" / "company_year_cross_validation_summary.csv",
    ]:
        assert_csv_has_header(path)
    check_human_validation(package_root)
    check_banned_and_paths(package_root)
    print("score_package_checks: ok")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except AssertionError as exc:
        print(f"score_package_checks: FAIL: {exc}", file=sys.stderr)
        raise SystemExit(1)
